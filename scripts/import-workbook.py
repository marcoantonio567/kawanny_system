"""Read-only extraction. Keeps original row references; never changes the workbook."""
import openpyxl, json, re, unicodedata, hashlib, sys
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

root = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1]) if len(sys.argv)>1 else root.parent / 'STUDIO KS.xlsx'
w = openpyxl.load_workbook(source, read_only=True, data_only=True)
months = ['JAN','FEV','MAR','ABR','MAI','JUN','JUL','AGO','SET','OUT','NOV','DEZ']
out = {k:[] for k in ['clients','services','categories','movements','goals','reserves','invoices','issues']}
def clean(v): return re.sub(r'\s+',' ',str(v or '')).strip()
def key(v): return unicodedata.normalize('NFKD',clean(v)).encode('ascii','ignore').decode().lower()
def cents(v):
    if v is None or isinstance(v,bool): return None
    try: return int((Decimal(str(v))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except: return None
def issue(sheet,row,reason,raw):
    out['issues'].append(dict(id=f'{sheet}-{row}-{len(out["issues"])}',source=f'{sheet}!{row}',reason=reason,raw=json.dumps(raw,ensure_ascii=False,default=str)))
clients={}; services={}; cats={}
def client(name,phone='',address='',birthday=''):
    name=clean(name)
    if not name:return None
    k=key(name)
    if k not in clients:
        clients[k]=dict(id=f'c{len(clients)+1}',name=name,phone=clean(phone),address=clean(address),birthday=birthday,notes='Importado da planilha')
    return clients[k]['id']
def service(name,price=None):
    name=clean(name)
    if not name:return None
    k=key(name)
    if k not in services:services[k]=dict(id=f's{len(services)+1}',name=name,price_cents=price or 0,duration=60,active=1)
    return services[k]['id']
def category(name,kind):
    name=clean(name) or 'Sem categoria';k=(key(name),kind)
    if k not in cats: cats[k]=dict(id=f'cat{len(cats)+1}',name=name,kind=kind)
    return cats[k]['id']
for row in w['CLIENTES'].iter_rows(min_row=9,max_col=7,values_only=True):
    if row[1]:
        b=row[4]; birthday=b.strftime('%m-%d') if isinstance(b,(datetime,date)) else clean(b)
        client(row[1],row[2],row[3],birthday)
config=list(w['CONFIGURAÇÕES'].iter_rows(max_row=70,max_col=56,values_only=True))
prices=list(w['PAINEL DE CONTROLE (NÃO APAGUE)'].iter_rows(min_row=7,max_row=70,max_col=11,values_only=True))
price_map={key(r[5]):cents(r[10]) for r in prices if r[5] and isinstance(r[10],(int,float))}
for r in config[8:64]:
    if r[19]: service(r[19],price_map.get(key(r[19])))
for start,end,kind in [(10,21,'income'),(23,35,'expense'),(55,66,'expense')]:
    for r in config[start-1:end]:
        if r[7]:category(r[7],kind)
summaries=[]
for month,n in zip(months,range(1,13)):
    s=w[month];rows=list(s.iter_rows(max_col=42,values_only=True))
    get=lambda r,c: rows[r-1][c-1] if r<=len(rows) else None
    out['goals'].append(dict(id=f'g{n}',period=f'2025-{n:02}',income_cents=cents(get(10,8)) or 0,expense_cents=cents(get(12,8)) or 0,reserve_cents=cents(get(13,8)) or 0))
    for i,r in enumerate(rows[85:],86):
        if not any(r[j] is not None for j in [1,3,6,10,18,26]):continue
        amount=cents(r[26]);kind={'entrada':'income','saida':'expense'}.get(key(r[3]))
        raw={openpyxl.utils.get_column_letter(j+1):v for j,v in enumerate(r) if v is not None and v is not False}
        if not kind or amount is None or amount<=0:
            issue(month,i,'Tipo ou valor ausente/inválido; não contabilizado.',raw);continue
        if not isinstance(r[1],(datetime,date)):
            issue(month,i,'Data inválida; lançamento aguardando revisão.',raw);continue
        d=r[1].strftime('%Y-%m-%d')
        if r[1].year!=2025 or r[1].month!=n:
            issue(month,i,'Data difere do mês/ano da aba. Data original preservada.',raw)
        cat=category(r[6],kind);sid=service(r[10]);cid=client(r[18]);payment=clean((r[29] or r[34]) if kind=='income' else (r[34] or r[29])) or 'Não informado'
        out['movements'].append(dict(id=f'm-{month}-{i}',date=d,kind=kind,category_id=cat,service_id=sid,client_id=cid,description=clean(r[10] or r[6]) or 'Lançamento importado',amount_cents=amount,payment=payment,status='pending' if key(payment)=='fiado' else 'paid',installments=max(1,int(r[38] or 1)),source=f'{month}!{i}',notes='Próximo mês marcado na origem' if r[41] is True else ''))
    for i in range(50,55):
        if clean(get(i,2)) and (cents(get(i,12)) or 0)>0:
            rd=get(i,6)
            out['reserves'].append(dict(id=f'r{n}-{i}',period=f'2025-{n:02}',name=clean(get(i,2)),date=rd.strftime('%Y-%m-%d') if isinstance(rd,(datetime,date)) else f'2025-{n:02}-01',amount_cents=cents(get(i,12)),target_cents=cents(get(i,9)) or 0))
    for i in range(60,65):
        if clean(get(i,2)):
            out['invoices'].append(dict(id=f'i{n}-{i}',period=f'2025-{n:02}',name=clean(get(i,2)),previous_cents=cents(get(i,5)) or 0,amount_cents=cents(get(i,9)) or 0,paid=int(get(i,12) is True)))
    detail_income=sum(x['amount_cents'] for x in out['movements'] if x['source'].startswith(month+'!') and x['kind']=='income')
    detail_expense=sum(x['amount_cents'] for x in out['movements'] if x['source'].startswith(month+'!') and x['kind']=='expense')
    cached=cents(get(10,11))
    if cached is not None and abs(detail_income-cached)>1:issue(month,10,'Resumo de entradas difere dos lançamentos. O sistema usa os registros válidos.',{'resumo_centavos':cached,'lancamentos_centavos':detail_income})
    summaries.append(dict(month=month,income=detail_income,expense=detail_expense,cached=cached))
out['clients']=list(clients.values());out['services']=list(services.values());out['categories']=list(cats.values())
out['metadata']=dict(filename=source.name,sha256=hashlib.sha256(source.read_bytes()).hexdigest(),summaries=summaries)
(root/'data').mkdir(exist_ok=True)
(root/'data'/'seed.json').write_text(json.dumps(out,ensure_ascii=False,default=str,separators=(',',':')),encoding='utf8')
print(json.dumps({'counts':{k:len(v) for k,v in out.items() if isinstance(v,list)},'summaries':summaries},ensure_ascii=False))
